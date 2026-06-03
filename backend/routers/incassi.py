from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from typing import List, Dict, Any, Optional
from datetime import date
from io import BytesIO, StringIO
import csv
import re

from motor.motor_asyncio import AsyncIOMotorDatabase

from models.incasso import MovimentoIncasso, MatchIncasso, ImportIncassiPreview, ConfermaIncassoRequest, ImportIncassiLog
from models.rata import StatoRata
from models.user import UserInDB, UserRole
from routers.auth import get_current_user, get_db
from services.audit import log_audit

router = APIRouter(prefix="/incassi", tags=["Parser Incassi"])


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _parse_amount(value: Any) -> float:
    raw = _norm(value).replace("€", "").replace(" ", "")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    else:
        raw = raw.replace(",", ".")
    return float(raw)


def _parse_date(value: Any) -> Optional[date]:
    raw = _norm(value)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            from datetime import datetime
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    return None


def _pick(row: Dict[str, Any], names: List[str]) -> Any:
    lowered = {str(k).lower().strip(): v for k, v in row.items()}
    for name in names:
        if name in lowered:
            return lowered[name]
    return None


def _rows_from_csv(content: bytes) -> List[Dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="ignore")
    sample = text[:2048]
    dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
    reader = csv.DictReader(StringIO(text), dialect=dialect)
    return [dict(r) for r in reader]


def _rows_from_xlsx(content: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    parsed = []
    for values in rows[1:]:
        parsed.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    return parsed


def _movimento_from_row(row: Dict[str, Any]) -> Optional[MovimentoIncasso]:
    amount_raw = _pick(row, ["importo", "amount", "accredito", "entrate", "dare", "avere"])
    causale = _pick(row, ["causale", "descrizione", "description", "operazione", "note"])
    if amount_raw is None or not causale:
        return None
    try:
        amount = _parse_amount(amount_raw)
    except Exception:
        return None
    if amount <= 0:
        return None
    return MovimentoIncasso(
        data_movimento=_parse_date(_pick(row, ["data", "data movimento", "date", "contabile"])),
        importo=amount,
        causale=_norm(causale),
        ordinante=_norm(_pick(row, ["ordinante", "mittente", "payer", "beneficiario"])),
        riferimento=_norm(_pick(row, ["riferimento", "cro", "trn", "id operazione"])),
        raw=row,
    )


async def _enrich_rata(db: AsyncIOMotorDatabase, rata: Dict[str, Any]) -> Dict[str, Any]:
    enriched = dict(rata)
    contratto = await db.contratti.find_one({"id": rata.get("contratto_id")}, {"_id": 0})
    if contratto:
        enriched["contratto_codice"] = contratto.get("codice_contratto")
        soggetto = await db.soggetti.find_one({"id": contratto.get("affittuario_id")}, {"_id": 0, "nome": 1})
        enriched["affittuario_nome"] = soggetto.get("nome") if soggetto else None
    return enriched


def _movimento_key(mov: MovimentoIncasso) -> str:
    return "|".join([
        (mov.data_movimento.isoformat() if mov.data_movimento else ""),
        f"{mov.importo:.2f}",
        (mov.riferimento or "").lower(),
        re.sub(r"\s+", " ", mov.causale.lower()).strip()[:120],
    ])


def _score_match(mov: MovimentoIncasso, rata: Dict[str, Any]) -> tuple[int, List[str]]:
    score = 0
    motivi = []
    causale = f"{mov.causale} {mov.ordinante or ''}".lower()
    importo_rata = float(rata.get("importo", 0))

    if abs(mov.importo - importo_rata) < 0.01:
        score += 45
        motivi.append("importo esatto")
    elif 0 < mov.importo < importo_rata:
        score += 20
        motivi.append("importo parziale compatibile")

    codice = str(rata.get("contratto_codice") or "").lower()
    if codice and codice in causale:
        score += 30
        motivi.append("codice contratto in causale")

    affittuario = str(rata.get("affittuario_nome") or "").lower()
    tokens = [t for t in re.split(r"\W+", affittuario) if len(t) > 2]
    matching_tokens = [t for t in tokens if t in causale]
    if matching_tokens:
        score += min(25, 10 * len(matching_tokens))
        motivi.append("nome affittuario in causale")

    periodo = str(rata.get("periodo") or "")
    if periodo and periodo in causale:
        score += 15
        motivi.append("periodo in causale")

    return min(score, 100), motivi


@router.post("/import-preview", response_model=ImportIncassiPreview)
async def import_preview(
    file: UploadFile = File(...),
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db),
):
    if current_user.ruolo == UserRole.LETTURA:
        raise HTTPException(status_code=403, detail="Permessi insufficienti")

    content = await file.read()
    suffix = (file.filename or "").lower().split(".")[-1]
    if suffix == "csv":
        rows = _rows_from_csv(content)
    elif suffix in {"xlsx", "xlsm"}:
        rows = _rows_from_xlsx(content)
    else:
        raise HTTPException(status_code=400, detail="Sono supportati file CSV e XLSX")

    movimenti = [m for row in rows if (m := _movimento_from_row(row))]
    pending = await db.rate.find(
        {"stato": {"$in": [StatoRata.DA_INCASSARE.value, StatoRata.IN_RITARDO.value, StatoRata.PARZIALE.value]}},
        {"_id": 0}
    ).to_list(5000)
    enriched_pending = [await _enrich_rata(db, r) for r in pending]

    matches = []
    duplicate_count = 0
    for mov in movimenti:
        mov_key = _movimento_key(mov)
        duplicato = await db.incassi_movimenti.find_one({"key": mov_key}, {"_id": 0, "id": 1})
        if duplicato:
            duplicate_count += 1

        best = None
        best_score = 0
        best_motivi = []
        for rata in enriched_pending:
            score, motivi = _score_match(mov, rata)
            if score > best_score:
                best, best_score, best_motivi = rata, score, motivi

        if best_score >= 80:
            stato = "match_certo"
        elif best_score >= 50:
            stato = "match_probabile"
        elif best_score > 0:
            stato = "da_revisionare"
        else:
            stato = "non_abbinato"

        matches.append(MatchIncasso(
            movimento=mov,
            rata_id=best.get("id") if best else None,
            score=best_score,
            motivi=best_motivi,
            stato_match=stato,
            rata=best,
            duplicato=bool(duplicato),
        ))

    log = ImportIncassiLog(
        filename=file.filename or "movimenti",
        total_movimenti=len(movimenti),
        matched_count=len([m for m in matches if m.stato_match in ("match_certo", "match_probabile")]),
        duplicate_count=duplicate_count,
        uploaded_by=current_user.id,
    )
    log_doc = log.model_dump()
    log_doc["created_at"] = log_doc["created_at"].isoformat()
    await db.import_incassi_log.insert_one(log_doc)

    return ImportIncassiPreview(matches=matches, total_movimenti=len(movimenti))


@router.post("/conferma")
async def conferma_incasso(
    payload: ConfermaIncassoRequest,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db),
):
    if current_user.ruolo == UserRole.LETTURA:
        raise HTTPException(status_code=403, detail="Permessi insufficienti")

    existing = await db.rate.find_one({"id": payload.rata_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Rata non trovata")

    importo_rata = float(existing.get("importo", 0))
    stato = StatoRata.INCASSATO.value if payload.importo >= importo_rata else StatoRata.PARZIALE.value
    update = {
        "stato": stato,
        "data_incasso": (payload.data_incasso or date.today()).isoformat(),
        "metodo": payload.metodo,
        "riferimento": payload.riferimento,
        "note": payload.note,
    }
    if stato == StatoRata.PARZIALE.value:
        update["importo_parziale"] = payload.importo

    await db.rate.update_one({"id": payload.rata_id}, {"$set": update})
    mov_key = "|".join([
        update["data_incasso"],
        f"{payload.importo:.2f}",
        (payload.riferimento or "").lower(),
        (payload.note or "").lower()[:120],
    ])
    await db.incassi_movimenti.update_one(
        {"key": mov_key},
        {"$set": {
            "key": mov_key,
            "rata_id": payload.rata_id,
            "importo": payload.importo,
            "data_incasso": update["data_incasso"],
            "riferimento": payload.riferimento,
            "created_by": current_user.id,
        }},
        upsert=True,
    )
    await log_audit(db, "rate", payload.rata_id, "incasso_importato", existing, update, current_user.id)
    return {"message": "Incasso registrato", "stato": stato}


@router.get("/rate-pendenti")
async def rate_pendenti(
    limit: int = 500,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db),
):
    rate = await db.rate.find(
        {"stato": {"$in": [StatoRata.DA_INCASSARE.value, StatoRata.IN_RITARDO.value, StatoRata.PARZIALE.value]}},
        {"_id": 0}
    ).sort("data_scadenza", 1).limit(limit).to_list(limit)
    return [await _enrich_rata(db, r) for r in rate]


@router.get("/import-log", response_model=List[ImportIncassiLog])
async def import_log(
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db),
):
    return await db.import_incassi_log.find({}, {"_id": 0}).sort("created_at", -1).limit(50).to_list(50)
